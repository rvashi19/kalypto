# ruff: noqa: E501
"""Parse XLSX/CSV shipment sheets and map columns to export document fields."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

# Column alias map: canonical_field -> list of accepted header variants (lowercase, stripped)
COLUMN_ALIASES: dict[str, list[str]] = {
    "item_number": ["item no", "item number", "sl no", "sr no", "s.no", "#", "sr.no"],
    "product_description": ["product", "item", "description", "goods description", "commodity", "particulars", "product description", "item description", "goods"],
    "hsn_code": ["hs code", "hsn", "hsn code", "itc hs", "itc-hs", "hs", "tariff code", "hsn/sac"],
    "quantity": ["qty", "quantity", "nos", "no of units", "units"],
    "unit": ["uom", "unit", "unit of measure", "uom/unit", "measure"],
    "unit_price": ["price", "rate", "unit price", "unit rate", "price per unit", "rate per unit"],
    "total_value": ["amount", "value", "total", "total value", "total amount", "line total", "net amount"],
    "net_weight": ["net wt", "net weight", "nett weight", "net wt (kg)", "net weight (kg)"],
    "gross_weight": ["gross wt", "gross weight", "gross wt (kg)", "gross weight (kg)"],
    "package_count": ["cartons", "boxes", "packages", "no of packages", "no. of packages", "pkgs", "ctns", "bags", "drums", "pcs"],
    "package_type": ["package type", "packing", "packing type", "carton type", "pkg type"],
    "dimensions": ["dimensions", "dimension", "size", "l x w x h", "l*w*h"],
    "batch_lot_number": ["batch", "lot", "batch no", "lot no", "batch number", "lot number"],
    "country_of_origin": ["origin", "country of origin", "made in"],
}

REQUIRED_ITEM_FIELDS = ["product_description", "hsn_code", "quantity", "unit_price", "total_value"]
REQUIRED_PACK_FIELDS = ["net_weight", "gross_weight", "package_count"]


def _normalize_header(raw: str) -> str:
    return re.sub(r"\s+", " ", raw.strip().lower())


def _match_column(header: str) -> str | None:
    norm = _normalize_header(header)
    for field, aliases in COLUMN_ALIASES.items():
        if norm in aliases or norm == field.replace("_", " "):
            return field
    return None


def detect_column_mapping(headers: list[str]) -> tuple[dict[str, str], list[str]]:
    """Return (mapping: raw_header -> field, unmatched_headers)."""
    mapping: dict[str, str] = {}
    matched_fields: set[str] = set()
    unmatched: list[str] = []
    for h in headers:
        field = _match_column(h)
        if field and field not in matched_fields:
            mapping[h] = field
            matched_fields.add(field)
        else:
            unmatched.append(h)
    return mapping, unmatched


def _parse_float(val: Any) -> float | None:
    if val is None or str(val).strip() in ("", "-", "N/A", "n/a"):
        return None
    try:
        return float(str(val).replace(",", "").strip())
    except (ValueError, TypeError):
        return None


def _parse_int(val: Any) -> int | None:
    f = _parse_float(val)
    return int(f) if f is not None else None


def parse_rows(rows: list[dict[str, Any]], mapping: dict[str, str]) -> list[dict[str, Any]]:
    """Apply column mapping to raw rows, returning normalized item dicts."""
    # Invert mapping: field -> raw_header
    field_to_header = {v: k for k, v in mapping.items()}
    items: list[dict[str, Any]] = []
    for i, row in enumerate(rows):
        def _get(field: str) -> Any:
            h = field_to_header.get(field)
            return row.get(h) if h else None

        item: dict[str, Any] = {
            "row_index": i + 2,  # 1-based + header row
            "item_number": _parse_int(_get("item_number")) or (i + 1),
            "product_description": str(_get("product_description") or "").strip() or None,
            "hsn_code": str(_get("hsn_code") or "").strip().replace(" ", "") or None,
            "quantity": _parse_float(_get("quantity")),
            "unit": str(_get("unit") or "").strip() or None,
            "unit_price": _parse_float(_get("unit_price")),
            "total_value": _parse_float(_get("total_value")),
            "net_weight": _parse_float(_get("net_weight")),
            "gross_weight": _parse_float(_get("gross_weight")),
            "package_count": _parse_int(_get("package_count")),
            "package_type": str(_get("package_type") or "").strip() or None,
            "dimensions": str(_get("dimensions") or "").strip() or None,
            "batch_lot_number": str(_get("batch_lot_number") or "").strip() or None,
            "country_of_origin": str(_get("country_of_origin") or "").strip() or None,
        }
        # Skip completely empty rows
        if not any(v for k, v in item.items() if k not in ("row_index", "item_number")):
            continue
        items.append(item)
    return items


def detect_missing_fields(items: list[dict[str, Any]]) -> list[str]:
    """Return list of field names missing from at least one item."""
    missing: set[str] = set()
    check = REQUIRED_ITEM_FIELDS + REQUIRED_PACK_FIELDS
    for item in items:
        for field in check:
            if item.get(field) is None:
                missing.add(field)
    return sorted(missing)


def parse_xlsx(file_bytes: bytes, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, Any]], list[str]]:
    """Parse XLSX bytes. Returns (headers, rows_as_dicts, sheet_names)."""
    import openpyxl  # already in deps

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet_names: list[str] = wb.sheetnames
    ws = wb[sheet_name] if sheet_name and sheet_name in sheet_names else wb.active
    rows_iter = list(ws.iter_rows(values_only=True))
    if not rows_iter:
        return [], [], sheet_names

    # Find header row (first non-empty row)
    headers: list[str] = []
    data_start = 0
    for idx, row in enumerate(rows_iter):
        cells = [str(c).strip() if c is not None else "" for c in row]
        non_empty = [c for c in cells if c]
        if len(non_empty) >= 2:
            headers = cells
            data_start = idx + 1
            break

    rows: list[dict[str, Any]] = []
    for row in rows_iter[data_start:]:
        row_dict = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
        rows.append(row_dict)

    return headers, rows, sheet_names


def parse_csv(file_bytes: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    """Parse CSV bytes. Returns (headers, rows_as_dicts)."""
    text = file_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    headers = list(reader.fieldnames or [])
    rows = list(reader)
    return headers, rows
