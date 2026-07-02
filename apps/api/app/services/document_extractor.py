# ruff: noqa: E501
"""Extract structured shipment data from any uploaded document.

Supported input formats:
  PDF   — pdfplumber text extraction → LLM parsing
  DOCX  — python-docx text extraction → LLM parsing
  XLSX  — openpyxl → existing column-mapper (no LLM needed)
  CSV   — csv.DictReader → existing column-mapper (no LLM needed)

The LLM extraction prompt asks the model to return a JSON object matching
the ShipmentPayload schema. Fields not found in the document are null.
The response also includes a `missing_fields` list and `confidence` score.
"""

from __future__ import annotations

import importlib
import io
import json
import re
from typing import Any

_EXTRACTION_SYSTEM = """You are an expert export document parser for Indian exporters.

Extract ALL available information from the provided document text and return it as strict JSON.

Return exactly this structure (use null for any field not found):
{
  "exporter": {
    "company_name": null,
    "address": null,
    "city": null,
    "state": null,
    "postal_code": null,
    "country": null,
    "iec": null,
    "gstin": null,
    "email": null,
    "bank_name": null,
    "bank_account": null,
    "ifsc_swift": null,
    "ad_code": null
  },
  "buyer": {
    "buyer_name": null,
    "buyer_address": null,
    "buyer_country": null,
    "consignee_name": null,
    "consignee_address": null,
    "notify_party": null
  },
  "shipment": {
    "invoice_number": null,
    "invoice_date": null,
    "proforma_invoice_number": null,
    "buyer_order_number": null,
    "country_of_origin": null,
    "country_of_final_destination": null,
    "port_of_loading": null,
    "port_of_discharge": null,
    "incoterm": null,
    "mode_of_transport": null,
    "currency": null,
    "payment_terms": null,
    "marks_and_numbers": null,
    "container_number": null,
    "vessel_flight_number": null,
    "expected_shipment_date": null
  },
  "items": [
    {
      "item_number": null,
      "product_description": null,
      "hsn_code": null,
      "quantity": null,
      "unit": null,
      "unit_price": null,
      "total_value": null,
      "net_weight": null,
      "gross_weight": null,
      "package_count": null,
      "package_type": null
    }
  ],
  "packing": {
    "total_packages": null,
    "package_type": null,
    "total_net_weight": null,
    "total_gross_weight": null,
    "freight": null,
    "insurance": null
  },
  "declarations": {
    "authorized_signatory_name": null,
    "authorized_signatory_designation": null,
    "place_of_issue": null,
    "date_of_issue": null
  },
  "confidence": 0,
  "missing_fields": [],
  "notes": null
}

Rules:
- Dates: use YYYY-MM-DD format if possible
- Numbers: return as numbers (not strings) for quantity, unit_price, total_value, weights, package_count
- HSN codes: strip spaces, digits only
- Incoterms: 3-letter code (EXW, FOB, CIF, etc.)
- currency: 3-letter ISO code (USD, EUR, INR, etc.)
- confidence: 0-100 based on how much of the document you could parse
- missing_fields: list field paths that were not found, e.g. ["exporter.iec", "shipment.port_of_loading"]
- notes: any important caveats or ambiguities (or null)
- Return STRICT JSON only — no markdown, no explanation outside the JSON
"""


# ── Text extraction ──────────────────────────────────────────────────────────

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract all text from a PDF using pdfplumber."""
    import pdfplumber  # already in requirements

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        pages = []
        for page in pdf.pages:
            # Extract text
            text = page.extract_text() or ""
            # Also extract tables and append as TSV-ish text
            for table in (page.extract_tables() or []):
                for row in table:
                    cells = [str(c or "").strip() for c in row]
                    if any(cells):
                        text += "\n" + "\t".join(cells)
            pages.append(text)
    return "\n\n--- PAGE BREAK ---\n\n".join(pages)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract all text from a DOCX file using python-docx."""
    import docx  # python-docx

    doc = docx.Document(io.BytesIO(file_bytes))
    lines: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append("\t".join(cells))

    return "\n".join(lines)


def extract_text_from_xlsx_as_text(file_bytes: bytes, sheet_name: str | None = None) -> str:
    """Convert XLSX to a plain-text table for LLM parsing (used when column mapping fails)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    lines: list[str] = []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if any(cells):
            lines.append("\t".join(cells))
    return "\n".join(lines)


def extract_text_from_csv(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8-sig", errors="replace")


# ── LLM provider (mirrors hsn_llm_verify pattern) ───────────────────────────

def _provider() -> tuple[str, str, str | None, str] | None:
    from app.core.settings import get_settings
    settings = get_settings()
    if settings.openai_api_key:
        return ("openai", settings.openai_api_key, None, settings.openai_model)
    if settings.xai_api_key:
        return ("xai", settings.xai_api_key, settings.xai_base_url, settings.xai_model)
    if settings.groq_api_key:
        return ("groq", settings.groq_api_key, "https://api.groq.com/openai/v1", settings.groq_model)
    return None


def is_llm_configured() -> bool:
    return _provider() is not None


def _call_llm(document_text: str) -> dict[str, Any]:
    provider = _provider()
    if provider is None:
        raise RuntimeError("No LLM provider configured. Set OPENAI_API_KEY, XAI_API_KEY, or GROQ_API_KEY.")

    name, api_key, base_url, model = provider

    # Truncate to ~12k chars to stay within token limits
    truncated = document_text[:12000]
    if len(document_text) > 12000:
        truncated += "\n\n[... document truncated for length ...]"

    user_message = f"Extract all export document fields from the following document:\n\n{truncated}"

    # Groq only supports the Chat Completions API — route through the dedicated
    # native-SDK client (JSON mode) rather than OpenAI's Responses API.
    if name == "groq":
        from app.services.groq_client import GroqClientError, call_groq
        try:
            return call_groq(system_prompt=_EXTRACTION_SYSTEM, user_message=user_message)
        except GroqClientError as e:
            raise RuntimeError(f"LLM extraction failed: {e}") from e

    # OpenAI / xAI: Chat Completions with JSON mode is universally supported.
    openai_module: Any = importlib.import_module("openai")
    client = (
        openai_module.OpenAI(api_key=api_key, base_url=base_url)
        if base_url
        else openai_module.OpenAI(api_key=api_key)
    )

    try:
        completion = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": _EXTRACTION_SYSTEM},
                {"role": "user", "content": user_message},
            ],
            temperature=0.2,
            response_format={"type": "json_object"},
        )
    except Exception as e:
        raise RuntimeError(f"LLM extraction failed: {e}") from e

    text = (completion.choices[0].message.content or "").strip()
    text = text.removeprefix("```json").removeprefix("```").removesuffix("```").strip()

    # Find JSON object in response even if there's surrounding text
    match = re.search(r"\{[\s\S]*\}", text)
    if not match:
        raise RuntimeError(f"LLM did not return valid JSON. Response: {text[:200]}")

    try:
        return json.loads(match.group())
    except json.JSONDecodeError as e:
        raise RuntimeError(f"LLM returned malformed JSON: {e}") from e


# ── Public extraction API ────────────────────────────────────────────────────

_TEXT_EXTRACTORS = {
    "pdf": extract_text_from_pdf,
    "docx": extract_text_from_docx,
    "doc": extract_text_from_docx,
    "csv": extract_text_from_csv,
}

SUPPORTED_EXTRACT_EXTENSIONS = {"pdf", "docx", "doc", "xlsx", "xls", "csv"}


def detect_file_type(filename: str, content_type: str) -> str:
    """Return lowercase extension. Raises ValueError for unsupported types."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in SUPPORTED_EXTRACT_EXTENSIONS:
        return ext
    # Fallback to content-type sniffing
    ct = content_type.lower()
    if "pdf" in ct:
        return "pdf"
    if "spreadsheet" in ct or "excel" in ct:
        return "xlsx"
    if "csv" in ct or "text/plain" in ct:
        return "csv"
    if "wordprocessingml" in ct or "msword" in ct:
        return "docx"
    raise ValueError(f"Unsupported file type: {filename} ({content_type})")


def extract_and_parse(
    file_bytes: bytes,
    filename: str,
    content_type: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """
    Main entry point. Returns a dict with:
      {
        "extracted": <ShipmentPayload-shaped dict, nulls where not found>,
        "confidence": 0-100,
        "missing_fields": ["exporter.iec", ...],
        "notes": str | null,
        "file_type": "pdf" | "xlsx" | "csv" | "docx",
        "used_llm": bool,
      }
    """
    file_type = detect_file_type(filename, content_type)

    # For XLSX/XLS: try column-mapper first, fall through to LLM if too sparse
    if file_type in ("xlsx", "xls"):
        from app.services.document_sheet_parser import (
            detect_column_mapping,
            detect_missing_fields,
            parse_rows,
            parse_xlsx,
        )
        headers, rows, sheet_names = parse_xlsx(file_bytes, sheet_name)
        mapping, unmatched = detect_column_mapping(headers)

        if mapping and rows:
            parsed_items = parse_rows(rows, mapping)
            missing = detect_missing_fields(parsed_items)
            # Build a partial payload from spreadsheet items
            payload = _items_to_payload(parsed_items)
            return {
                "extracted": payload,
                "confidence": max(20, 80 - len(missing) * 10),
                "missing_fields": _payload_missing_fields(payload) + [f"items[*].{f}" for f in missing],
                "notes": f"Parsed from spreadsheet. {len(unmatched)} unrecognised columns ignored." if unmatched else None,
                "file_type": file_type,
                "used_llm": False,
                "sheet_names": sheet_names,
            }
        # Fall through to LLM with spreadsheet text
        document_text = extract_text_from_xlsx_as_text(file_bytes, sheet_name)

    elif file_type == "csv":
        document_text = extract_text_from_csv(file_bytes)
        # Try column-mapper first
        from app.services.document_sheet_parser import (
            detect_column_mapping, detect_missing_fields, parse_rows, parse_csv,
        )
        headers, rows = parse_csv(file_bytes)
        mapping, _ = detect_column_mapping(headers)
        if mapping and rows:
            parsed_items = parse_rows(rows, mapping)
            payload = _items_to_payload(parsed_items)
            return {
                "extracted": payload,
                "confidence": 65,
                "missing_fields": _payload_missing_fields(payload),
                "notes": "Parsed from CSV. Exporter, buyer, and shipment details not found — please fill those in.",
                "file_type": file_type,
                "used_llm": False,
                "sheet_names": [],
            }

    else:
        extractor = _TEXT_EXTRACTORS[file_type]
        document_text = extractor(file_bytes)

    # LLM extraction
    result = _call_llm(document_text)

    return {
        "extracted": {
            "exporter": result.get("exporter") or {},
            "buyer": result.get("buyer") or {},
            "shipment": result.get("shipment") or {},
            "items": result.get("items") or [],
            "packing": result.get("packing") or {},
            "declarations": result.get("declarations") or {},
        },
        "confidence": result.get("confidence", 50),
        "missing_fields": result.get("missing_fields") or [],
        "notes": result.get("notes"),
        "file_type": file_type,
        "used_llm": True,
        "sheet_names": [],
    }


# ── Helpers ──────────────────────────────────────────────────────────────────

def _items_to_payload(items: list[dict[str, Any]]) -> dict[str, Any]:
    """Wrap parsed spreadsheet items into a ShipmentPayload-shaped skeleton."""
    return {
        "exporter": {},
        "buyer": {},
        "shipment": {},
        "items": items,
        "packing": {
            "total_packages": sum(i.get("package_count") or 0 for i in items) or None,
            "total_net_weight": sum(i.get("net_weight") or 0 for i in items) or None,
            "total_gross_weight": sum(i.get("gross_weight") or 0 for i in items) or None,
        },
        "declarations": {},
    }


def _payload_missing_fields(payload: dict[str, Any]) -> list[str]:
    """Return dotted paths for top-level required fields that are empty."""
    missing: list[str] = []
    exp = payload.get("exporter") or {}
    if not exp.get("company_name"):
        missing.append("exporter.company_name")
    if not exp.get("iec"):
        missing.append("exporter.iec")
    byr = payload.get("buyer") or {}
    if not byr.get("buyer_name"):
        missing.append("buyer.buyer_name")
    if not byr.get("buyer_country"):
        missing.append("buyer.buyer_country")
    shp = payload.get("shipment") or {}
    for f in ("invoice_number", "incoterm", "currency", "port_of_loading", "port_of_discharge"):
        if not shp.get(f):
            missing.append(f"shipment.{f}")
    return missing
