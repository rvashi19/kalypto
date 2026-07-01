"""OCR and field extraction for uploaded shipment documents."""

from __future__ import annotations

import base64
import logging
import re
from pathlib import Path
from typing import Any

from app.services.groq_client import GroqClientError, call_groq

logger = logging.getLogger(__name__)

_FIELD_EXTRACT_SYSTEM = """You are an expert Indian export document reader.
Extract key fields from the raw text of a shipment document and return a JSON object.
Only include fields that are clearly present in the text. Return null for missing fields.

Return a JSON object with these fields (use null if absent):
{
  "document_number": null,
  "invoice_number": null,
  "document_date": null,
  "buyer_name": null,
  "seller_name": null,
  "hsn_code": null,
  "product_description": null,
  "quantity": null,
  "unit_of_measure": null,
  "net_weight_kg": null,
  "gross_weight_kg": null,
  "fob_value": null,
  "cif_value": null,
  "currency": null,
  "incoterm": null,
  "payment_term": null,
  "port_of_loading": null,
  "port_of_discharge": null,
  "vessel_flight_no": null,
  "bl_awb_number": null,
  "shipping_bill_no": null,
  "igst_amount": null,
  "gstin": null,
  "iec_code": null,
  "ad_code": null,
  "marks_and_numbers": null
}

Return ONLY the JSON object. No markdown."""


def _extract_text_from_pdf(file_path: str) -> str:
    try:
        import pdfplumber  # noqa: PLC0415
    except ImportError:
        return ""
    text_parts: list[str] = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages[:10]:
            t = page.extract_text()
            if t:
                text_parts.append(t)
    return "\n".join(text_parts)


def _extract_text_from_image(file_path: str) -> str:
    """Use Groq vision to extract text from an image document."""
    from groq import Groq  # noqa: PLC0415

    from app.core.settings import get_settings  # noqa: PLC0415

    settings = get_settings()
    if not settings.groq_api_key:
        return ""

    with open(file_path, "rb") as f:
        image_data = base64.b64encode(f.read()).decode("utf-8")

    suffix = Path(file_path).suffix.lower()
    media_type_map = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
    }
    media_type = media_type_map.get(suffix, "image/jpeg")

    client = Groq(api_key=settings.groq_api_key)
    try:
        completion = client.chat.completions.create(
            model="llama-3.2-11b-vision-preview",
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:{media_type};base64,{image_data}"},
                        },
                        {
                            "type": "text",
                            "text": (
                                "Extract all visible text from this export document. "
                                "Return just the raw text content, preserving structure."
                            ),
                        },
                    ],
                }
            ],
            temperature=0.1,
        )
        return completion.choices[0].message.content or ""
    except Exception as e:
        logger.warning("Groq vision extraction failed: %s", e)
        return ""


def _extract_text_from_excel(file_path: str) -> str:
    try:
        from openpyxl import load_workbook  # noqa: PLC0415
    except ImportError:
        return ""

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    text_parts: list[str] = []
    for worksheet in workbook.worksheets[:5]:
        text_parts.append(f"Sheet: {worksheet.title}")
        for row in worksheet.iter_rows(max_row=200, values_only=True):
            values = [str(value).strip() for value in row if value not in (None, "")]
            if values:
                if len(values) == 2:
                    text_parts.append(f"{values[0]}: {values[1]}")
                else:
                    text_parts.append(" | ".join(values))
    return "\n".join(text_parts)


def _first_match(patterns: list[str], text: str) -> str | None:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
        if match:
            return " ".join(match.group(1).strip().split())
    return None


def _extract_fields_deterministic(raw_text: str) -> dict[str, Any]:
    compact_text = re.sub(r"[ \t]+", " ", raw_text)
    fields: dict[str, Any] = {}
    pattern_map = {
        "invoice_number": [
            r"(?:invoice\s*(?:no\.?|number|#)\s*[:\-]?\s*)([A-Z0-9\-\/]{3,40})",
            r"(?:inv\s*(?:no\.?|#)\s*[:\-]?\s*)([A-Z0-9\-\/]{3,40})",
        ],
        "document_number": [
            r"(?:document\s*(?:no\.?|number|#)\s*[:\-]?\s*)([A-Z0-9\-\/]{3,40})",
        ],
        "document_date": [
            r"(?:date\s*[:\-]?\s*)(\d{1,2}[\/\-.]\d{1,2}[\/\-.]\d{2,4})",
            r"(?:date\s*[:\-]?\s*)(\d{4}[\/\-.]\d{1,2}[\/\-.]\d{1,2})",
        ],
        "buyer_name": [
            r"(?:buyer|consignee)\s*(?:name)?\s*[:\-]?\s*([A-Z0-9 &.,'\-]{3,120})",
        ],
        "seller_name": [
            r"(?:seller|exporter)\s*(?:name)?\s*[:\-]?\s*([A-Z0-9 &.,'\-]{3,120})",
        ],
        "hsn_code": [
            r"(?:hsn|hs code|itc\s*hs)\s*(?:code)?\s*[:\-]?\s*(\d{4,10})",
        ],
        "quantity": [
            r"(?:quantity|qty)\s*[:\-]?\s*([0-9,.]+)",
        ],
        "unit_of_measure": [
            r"(?:uom|unit(?:\s*of\s*measure)?)\s*[:\-]?\s*([A-Z]{2,20})",
        ],
        "net_weight_kg": [
            r"(?:net\s*weight|net\s*wt)\s*[:\-]?\s*([0-9,.]+)\s*(?:kg|kgs|kilograms)?",
        ],
        "gross_weight_kg": [
            r"(?:gross\s*weight|gross\s*wt)\s*[:\-]?\s*([0-9,.]+)\s*(?:kg|kgs|kilograms)?",
        ],
        "fob_value": [
            r"(?:fob\s*(?:value|amount)?)\s*[:\-]?\s*(?:INR|USD|EUR|GBP|AED)?\s*([0-9,.]+)",
        ],
        "cif_value": [
            r"(?:cif\s*(?:value|amount)?)\s*[:\-]?\s*(?:INR|USD|EUR|GBP|AED)?\s*([0-9,.]+)",
        ],
        "currency": [
            r"\b(INR|USD|EUR|GBP|AED|SGD)\b",
        ],
        "incoterm": [
            r"\b(EXW|FOB|CFR|CIF|DDP|DAP|FCA)\b",
        ],
        "payment_term": [
            r"(?:payment\s*terms?|terms\s*of\s*payment)\s*[:\-]?\s*([A-Z0-9 ,.\/\-]{3,120})",
        ],
        "port_of_loading": [
            r"(?:port\s*of\s*loading|pol)\s*[:\-]?\s*([A-Z0-9 ,.\/\-]{3,120})",
        ],
        "port_of_discharge": [
            r"(?:port\s*of\s*discharge|pod)\s*[:\-]?\s*([A-Z0-9 ,.\/\-]{3,120})",
        ],
        "bl_awb_number": [
            r"(?:bl|b\/l|awb|airway\s*bill)\s*(?:no\.?|number|#)?\s*[:\-]?\s*([A-Z0-9\-\/]{3,40})",
        ],
        "shipping_bill_no": [
            r"(?:shipping\s*bill)\s*(?:no\.?|number|#)?\s*[:\-]?\s*([A-Z0-9\-\/]{3,40})",
        ],
        "igst_amount": [
            r"(?:igst)\s*(?:amount|value)?\s*[:\-]?\s*(?:INR)?\s*([0-9,.]+)",
        ],
        "gstin": [
            r"\b([0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z])\b",
        ],
        "iec_code": [
            r"(?:iec)\s*(?:code|no\.?|number)?\s*[:\-]?\s*([A-Z0-9]{10})",
        ],
        "ad_code": [
            r"(?:ad\s*code)\s*[:\-]?\s*([0-9]{7,14})",
        ],
    }
    for field, patterns in pattern_map.items():
        value = _first_match(patterns, compact_text)
        if value:
            fields[field] = value

    product_match = _first_match(
        [
            (
                r"(?:description\s*of\s*goods|product\s*description|goods)"
                r"\s*[:\-]?\s*([A-Z0-9 ,.\/'\-]{3,160})"
            ),
        ],
        compact_text,
    )
    if product_match:
        fields["product_description"] = product_match

    if fields:
        fields["_extraction_method"] = "deterministic_text"
    return fields


def extract_fields(file_path: str, mime_type: str | None) -> dict[str, Any]:
    """Extract structured fields from a document file. Returns a dict of fields."""
    raw_text = ""

    if mime_type == "application/pdf":
        raw_text = _extract_text_from_pdf(file_path)
    elif mime_type in {"image/jpeg", "image/png", "image/webp"}:
        raw_text = _extract_text_from_image(file_path)
    elif mime_type in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    }:
        raw_text = _extract_text_from_excel(file_path)

    if not raw_text or len(raw_text.strip()) < 20:
        return {"_note": "Could not extract readable text from this document"}

    truncated = raw_text[:8000]
    deterministic_fields = _extract_fields_deterministic(truncated)

    try:
        fields = call_groq(
            system_prompt=_FIELD_EXTRACT_SYSTEM,
            user_message=f"Document text:\n{truncated}",
        )
        ai_fields = {k: v for k, v in fields.items() if v is not None}
        return {**deterministic_fields, **ai_fields, "_extraction_method": "ai_plus_text"}
    except GroqClientError as e:
        logger.warning("Field extraction Groq call failed: %s", e)
        if deterministic_fields:
            deterministic_fields["_note"] = (
                "AI extraction unavailable; deterministic text fields used."
            )
            return deterministic_fields
        return {"_note": f"AI extraction failed: {e}"}
